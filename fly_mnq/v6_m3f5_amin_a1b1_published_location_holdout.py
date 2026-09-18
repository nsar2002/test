from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

WORK=Path("v6_m3f5_work")
OUT=Path("v6_m3f5_out")
WORK.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)

SOURCE=WORK/"elife-56954-fig7-data1-v2.xlsx"
EXPECTED_SHA256="ba1d7c4dfaa2007aa3e71f668b72ee81845ea20094012f44450117a3a2165e9c"
SHEET="Fig 7"
SEED_A=2026091801
SEED_B=2026091802
N_BOOT=100000

def sha256(path:Path)->str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda:f.read(8<<20),b""):
            h.update(b)
    return h.hexdigest()

def numeric_triplet(values):
    out=[]
    for x in values:
        if isinstance(x,bool) or not isinstance(x,(int,float,np.integer,np.floating)):
            return None
        x=float(x)
        if not np.isfinite(x):
            return None
        out.append(x)
    return out

def extract_panel(ws,start_row,end_row):
    # Exact frozen C/H/V columns A:C only. Never access A4/B4 ranges.
    rows=[]
    for r in range(start_row,end_row+1):
        trip=numeric_triplet([ws.cell(r,1).value,ws.cell(r,2).value,ws.cell(r,3).value])
        if trip is not None:
            rows.append(trip)
    return np.asarray(rows,dtype=float)

def bootstrap_median(x,seed):
    rng=np.random.Generator(np.random.PCG64(seed))
    n=len(x)
    # Chunked to avoid unnecessary peak memory while preserving exact RNG stream.
    vals=np.empty(N_BOOT,dtype=float)
    chunk=10000
    pos=0
    while pos<N_BOOT:
        k=min(chunk,N_BOOT-pos)
        idx=rng.integers(0,n,size=(k,n),endpoint=False)
        vals[pos:pos+k]=np.median(x[idx],axis=1)
        pos+=k
    return [float(v) for v in np.quantile(vals,[0.025,0.975],method="linear")]

def summarize_a(a):
    c,h,v=a[:,0],a[:,1],a[:,2]
    d=h-0.5*(c+v)
    return {
        "complete_cases":int(len(a)),
        "composite_median":float(np.median(d)),
        "composite_ci95":bootstrap_median(d,SEED_A),
        "pairwise_median_H_minus_C":float(np.median(h-c)),
        "pairwise_median_H_minus_V":float(np.median(h-v)),
    }

def summarize_b(b):
    c,h,v=b[:,0],b[:,1],b[:,2]
    d=v-0.5*(c+h)
    return {
        "complete_cases":int(len(b)),
        "composite_median":float(np.median(d)),
        "composite_ci95":bootstrap_median(d,SEED_B),
        "pairwise_median_V_minus_C":float(np.median(v-c)),
        "pairwise_median_V_minus_H":float(np.median(v-h)),
    }

def main():
    got=sha256(SOURCE)
    if got!=EXPECTED_SHA256:
        raise RuntimeError(f"source hash mismatch {got}")

    wb=load_workbook(SOURCE,data_only=True,read_only=True)
    if SHEET not in wb.sheetnames:
        raise RuntimeError("frozen sheet missing")
    ws=wb[SHEET]

    # Verify only frozen A1/B1 labels. No adjacent response block is accessed.
    labels_a=[ws.cell(260,c).value for c in (1,2,3)]
    labels_b=[ws.cell(276,c).value for c in (1,2,3)]
    if [str(x).strip() for x in labels_a] != ["C","H","V"]:
        raise RuntimeError(f"A1 label mismatch {labels_a}")
    if [str(x).strip() for x in labels_b] != ["C","H","V"]:
        raise RuntimeError(f"B1 label mismatch {labels_b}")

    a=extract_panel(ws,261,274)
    b=extract_panel(ws,277,290)

    guardrails={
        "a4_b4_read":False,
        "c1_rescue_read":False,
        "source_pvalue_read":False,
        "voltage_to_gcamp_fit":False,
        "q_selected":False,
        "raw_per_neuron_values_emitted":False,
        "mnq_or_market_loaded":False,
    }

    if len(a)<8 or len(b)<8:
        report={
            "gate":"V6-M3F5_AMIN_A1B1_PUBLISHED_LOCATION_ORDINAL_HOLDOUT",
            "classification":"BLOCKED_M3F5_SOURCE_PAIRING_INSUFFICIENT",
            "source_sha256":got,
            "complete_cases":{"A1":int(len(a)),"B1":int(len(b))},
            "guardrails":guardrails,
        }
    else:
        sa=summarize_a(a)
        sb=summarize_b(b)
        passed=(
            sa["composite_ci95"][0]>0
            and sb["composite_ci95"][0]>0
            and sa["pairwise_median_H_minus_C"]>0
            and sa["pairwise_median_H_minus_V"]>0
            and sb["pairwise_median_V_minus_C"]>0
            and sb["pairwise_median_V_minus_H"]>0
        )
        cls="PASS_M3F5_INDEPENDENT_ORDINAL_APL_LOCALITY" if passed else "FAIL_M3F5_INDEPENDENT_ORDINAL_APL_LOCALITY"
        report={
            "gate":"V6-M3F5_AMIN_A1B1_PUBLISHED_LOCATION_ORDINAL_HOLDOUT",
            "classification":cls,
            "source_sha256":got,
            "sheet":SHEET,
            "frozen_ranges":{"A1":"A261:C274","B1":"A277:C290"},
            "bootstrap":{"replicates":N_BOOT,"generator":"PCG64","A1_seed":SEED_A,"B1_seed":SEED_B,"quantile_method":"linear"},
            "A1":sa,
            "B1":sb,
            "guardrails":guardrails,
        }

    (OUT/"v6_m3f5_result.json").write_text(json.dumps(report,indent=2,sort_keys=True)+"\n")
    # Intentionally emit summaries only; no row-level values.
    print(json.dumps(report,indent=2,sort_keys=True))

if __name__=="__main__":
    main()
