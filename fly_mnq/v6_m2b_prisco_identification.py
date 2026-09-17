from __future__ import annotations

import hashlib
import json
import math
import re
import subprocess
import urllib.parse
import urllib.request
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize_scalar

DOI = 'doi:10.5061/dryad.bk3j9kdd1'
API = 'https://datadryad.org/api/v2'
UA = 'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/152 Safari/537.36'
WORK = Path('v6_m2b_work')
OUT = Path('v6_m2b_out')
WORK.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)

FILES = {
    'apl_mch_oct': 'APL_GCaMP6m_Mch_vs_Oct_peaks_(Figure_2D).xlsx',
    'apl_dl_oct': 'APL_GCaMP6m_DL_vs_Oct_peaks_(Figure_2G).xlsx',
    'pn_mch_oct': 'GH146-GCaMP6m_Mch_vs_Oct_AL_average_activity_(Figure_3-figure_supplement_2B).xlsx',
    'pn_dl_oct': 'GH146-GCaMP6m_DL_vs_Oct_AL_average_activity_(Figure_3-figure_supplement_2D).xlsx',
    'g_delta': 'Inter-odour_delta_among_conditions_-_Figure_4-figure_supplement_1.xlsx',
    'locality_primary': 'APL_locality_PA_MP_vs_FA_MP_(Figure_5C).xlsx',
    'locality_secondary': 'APL_locality_PA_FA_vs_MP_FA_(Figure_5-figure_supplement_1B).xlsx',
    'mg_profiles': 'MB247-homer-GCaMP3_locality_of_MGs_responses_(Figure_5-figure_supplement_1A).xlsx',
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def request_json(url: str):
    req = urllib.request.Request(url, headers={'Accept':'application/json','X-API-Version':'2.1.0','User-Agent':UA})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.load(r)


def curl_json(url: str):
    cp = subprocess.run(['curl','--silent','--show-error','--location','--fail-with-body','--max-time','90','--retry','3','--retry-delay','1','--user-agent',UA,'--header','Accept: application/json',url], text=True, capture_output=True)
    if cp.returncode:
        raise RuntimeError(f'curl json failed exit={cp.returncode}: {cp.stderr[-300:]}')
    return json.loads(cp.stdout)


def curl_download(url: str, path: Path):
    cp = subprocess.run(['curl','--silent','--show-error','--location','--fail-with-body','--max-time','120','--retry','3','--retry-delay','1','--user-agent',UA,'--output',str(path),url], text=True, capture_output=True)
    if cp.returncode:
        if path.exists(): path.unlink()
        raise RuntimeError(f'curl file failed exit={cp.returncode}: {cp.stderr[-300:]}')


def embedded_list(obj: dict, preferred: str):
    emb = obj.get('_embedded', {})
    if preferred in emb and isinstance(emb[preferred], list): return emb[preferred]
    for v in emb.values():
        if isinstance(v, list) and (not v or isinstance(v[0], dict)): return v
    return []


def resolve_version():
    obj = request_json(f"{API}/datasets/{urllib.parse.quote(DOI, safe='')}")
    href = ((obj.get('_links') or {}).get('stash:version') or {}).get('href')
    m = re.search(r'/versions/(\d+)(?:$|[/?#])', href or '')
    if not m: raise RuntimeError('cannot resolve Dryad version id')
    return int(m.group(1)), obj


def resolve_files(version_id: int):
    url = f'{API}/versions/{version_id}/files?per_page=100'
    out=[]; seen=set()
    while url:
        if url in seen: raise RuntimeError('pagination loop')
        seen.add(url); obj=request_json(url); out.extend(embedded_list(obj,'stash:files'))
        href=((obj.get('_links') or {}).get('next') or {}).get('href')
        url = (href if href and href.startswith('http') else ('https://datadryad.org'+href if href else None))
    return out


def file_name(meta): return str(meta.get('path') or meta.get('name') or meta.get('fileName') or '').split('/')[-1]


def resolve_assembly(version_id: int):
    arr=curl_json(f'https://datadryad.org/downloads/zip_assembly_info/{version_id}.json')
    if not isinstance(arr,list): raise RuntimeError('assembly not list')
    out={}
    for x in arr:
        if not isinstance(x,dict): continue
        n=str(x.get('filename') or '').split('/')[-1]
        if n:
            if n in out: raise RuntimeError(f'duplicate assembly filename {n}')
            out[n]=x
    return out


def acquire_verified():
    version_id, rec = resolve_version()
    metas=resolve_files(version_id); assembly=resolve_assembly(version_id)
    by={}
    for m in metas: by.setdefault(file_name(m),[]).append(m)
    manifest=[]; paths={}
    for key,name in FILES.items():
        if len(by.get(name,[]))!=1 or name not in assembly: raise RuntimeError(f'exact source resolution failure: {name}')
        m=by[name][0]; dtype=str(m.get('digestType') or '').lower(); exp=str(m.get('digest') or '').lower(); size=int(m['size'])
        if dtype not in {'sha-256','sha256'} or not re.fullmatch(r'[0-9a-f]{64}',exp): raise RuntimeError(f'bad provider digest {name}')
        ax=assembly[name]; signed=ax.get('url'); asize=int(ax['size'])
        if not isinstance(signed,str) or not signed.startswith('https://') or asize!=size: raise RuntimeError(f'assembly identity failure {name}')
        p=WORK/name; curl_download(signed,p); got=sha256(p)
        if p.stat().st_size!=size or got!=exp: raise RuntimeError(f'byte identity failure {name}')
        manifest.append({'role':key,'filename':name,'bytes':size,'sha256':got})
        paths[key]=p
    return version_id, rec, manifest, paths


def ws(path: Path):
    return load_workbook(path, data_only=True, read_only=False).worksheets[0]


def finite_col(sheet, col: int, r1: int, r2: int):
    vals=[]
    for r in range(r1,r2+1):
        v=sheet.cell(r,col).value
        if isinstance(v,(int,float)) and np.isfinite(float(v)): vals.append(float(v))
    return np.asarray(vals,dtype=float)


def paired_diff(path: Path):
    s=ws(path); dif=[]
    for r in range(2,12):
        a=s.cell(r,1).value; b=s.cell(r,2).value
        if isinstance(a,(int,float)) and isinstance(b,(int,float)) and np.isfinite(float(a)) and np.isfinite(float(b)):
            dif.append(float(b)-float(a))
    if len(dif)<8: raise RuntimeError(f'paired directional source insufficient: {path.name} n={len(dif)}')
    med=float(np.median(np.asarray(dif)))
    return {'n':len(dif),'median_second_minus_first':med,'sign':int(np.sign(med)) if abs(med)>1e-12 else 0}


def identify_g(path: Path):
    s=ws(path); on=finite_col(s,3,2,11); off=finite_col(s,4,2,11)
    if len(on)<8 or len(off)<8:
        return {'status':'BLOCKED_M2B_G_SOURCE_INSUFFICIENT','n_on':len(on),'n_off':len(off)}
    don=float(np.median(on)); doff=float(np.median(off))
    if abs(doff)<=1e-12:
        return {'status':'FAIL_M2B_G_DENOMINATOR_ZERO','n_on':len(on),'n_off':len(off),'delta_on':don,'delta_off':doff}
    g=1.0-abs(don)/abs(doff)
    rng=np.random.Generator(np.random.PCG64(2026091703)); R=100000
    ion=rng.integers(0,len(on),size=(R,len(on))); ioff=rng.integers(0,len(off),size=(R,len(off)))
    bon=np.median(on[ion],axis=1); boff=np.median(off[ioff],axis=1)
    valid=np.abs(boff)>1e-12; valid_frac=float(np.mean(valid))
    bg=np.full(R,np.nan); bg[valid]=1.0-np.abs(bon[valid])/np.abs(boff[valid])
    finite=bg[np.isfinite(bg)]
    lo,hi=(float(np.quantile(finite,0.025,method='linear')),float(np.quantile(finite,0.975,method='linear'))) if finite.size else (math.nan,math.nan)
    passed=(abs(don)<abs(doff) and 0<g<1 and valid_frac>=0.995 and 0<lo<hi<1 and (hi-lo)<=0.25)
    return {'status':'PASS_G_IDENTIFIED' if passed else 'FAIL_G_UNDERIDENTIFIED','n_on':len(on),'n_off':len(off),'delta_on_median':don,'delta_off_median':doff,'g_apl_kc':g,'bootstrap_valid_fraction':valid_frac,'ci95':[lo,hi],'ci_width':hi-lo,'seed':2026091703,'replicates':R}


def objective_for(lam, SPA,SFA,SMP, opa, ofa, sempa, semfa):
    z=np.arange(5,dtype=float); K=np.exp(-np.abs(z[:,None]-z[None,:])/lam)
    apa=K@SPA; afa=K@SFA; amp=K@SMP
    if np.any(amp<=0): return math.inf
    p1=apa/amp; p2=afa/amp
    return float(np.sum(((opa-p1)/sempa)**2)+np.sum(((ofa-p2)/semfa)**2))


def slope(v):
    z=np.arange(5,dtype=float); return float(np.polyfit(z,np.asarray(v,dtype=float),1)[0])


def identify_lambda(primary: Path, secondary: Path, profiles: Path):
    sp=ws(profiles); SPA=finite_col(sp,1,3,7); SFA=finite_col(sp,4,3,7); SMP=finite_col(sp,7,3,7)
    sl=ws(primary); zlabels=finite_col(sl,1,3,7); opa=finite_col(sl,2,3,7); sempa=finite_col(sl,3,3,7); ofa=finite_col(sl,5,3,7); semfa=finite_col(sl,6,3,7)
    arrays=[SPA,SFA,SMP,zlabels,opa,sempa,ofa,semfa]
    if any(len(x)!=5 for x in arrays): return {'status':'BLOCKED_M2B_LAMBDA_SOURCE_INSUFFICIENT','lengths':[len(x) for x in arrays]}
    if len(np.unique(zlabels))!=5 or not (np.all(np.diff(zlabels)>0) or np.all(np.diff(zlabels)<0)):
        return {'status':'BLOCKED_M2B_LAMBDA_Z_LABELS_INVALID','z_labels':zlabels.tolist()}
    if any(np.any(x<=0) for x in [SPA,SFA,SMP,sempa,semfa]): return {'status':'BLOCKED_M2B_LAMBDA_NONPOSITIVE_SOURCE'}
    loB,hiB=.05,20.0
    res=minimize_scalar(lambda q: objective_for(math.exp(q),SPA,SFA,SMP,opa,ofa,sempa,semfa),bounds=(math.log(loB),math.log(hiB)),method='bounded',options={'xatol':1e-12})
    lopt=float(math.exp(res.x)); qopt=float(res.fun)
    grid=np.exp(np.linspace(math.log(loB),math.log(hiB),200001)); qgrid=np.empty_like(grid)
    for i,l in enumerate(grid): qgrid[i]=objective_for(float(l),SPA,SFA,SMP,opa,ofa,sempa,semfa)
    j=int(np.argmin(qgrid)); lg=float(grid[j]); qg=float(qgrid[j]); agree=abs(lopt-lg)/lg<=0.005
    mask=qgrid<=qg+3.841458820694124; inds=np.flatnonzero(mask)
    contiguous=bool(inds.size and np.all(np.diff(inds)==1)); touches=bool(inds.size and (inds[0]==0 or inds[-1]==len(grid)-1))
    llo=float(grid[inds[0]]) if inds.size else math.nan; lhi=float(grid[inds[-1]]) if inds.size else math.nan; width=lhi/llo if inds.size and llo>0 else math.inf
    primary_pass=bool(res.success and agree and loB<lg<hiB and contiguous and not touches and width<=1.50)
    z=np.arange(5,dtype=float); K=np.exp(-np.abs(z[:,None]-z[None,:])/lg); apa=K@SPA; afa=K@SFA; amp=K@SMP
    ss=ws(secondary); obs_pafa=finite_col(ss,2,3,7); obs_mpfa=finite_col(ss,5,3,7)
    sec_ok=False; sec={}
    if len(obs_pafa)==5 and len(obs_mpfa)==5 and np.all(afa>0):
        pred_pafa=apa/afa; pred_mpfa=amp/afa
        so1,sp1,so2,sp2=slope(obs_pafa),slope(pred_pafa),slope(obs_mpfa),slope(pred_mpfa)
        sec_ok=(all(abs(x)>1e-12 for x in [so1,sp1,so2,sp2]) and np.sign(so1)==np.sign(sp1) and np.sign(so2)==np.sign(sp2))
        sec={'observed_pa_fa_slope':so1,'predicted_pa_fa_slope':sp1,'observed_mp_fa_slope':so2,'predicted_mp_fa_slope':sp2,'pass':bool(sec_ok)}
    status='PASS_LAMBDA_IDENTIFIED' if primary_pass and sec_ok else ('FAIL_M2B_LOCALITY_SECONDARY_DIRECTION' if primary_pass else 'LAMBDA_Z_UNDERIDENTIFIED')
    return {'status':status,'lambda_z':lg,'q_min':qg,'optimizer_lambda':lopt,'optimizer_grid_relative_agreement':abs(lopt-lg)/lg,'profile_ci95':[llo,lhi],'profile_width_ratio':width,'profile_contiguous':contiguous,'profile_touches_boundary':touches,'secondary':sec,'z_labels_source':zlabels.tolist()}


def main():
    version_id,rec,manifest,paths=acquire_verified()
    ff={k:paired_diff(paths[k]) for k in ['apl_mch_oct','pn_mch_oct','apl_dl_oct','pn_dl_oct']}
    ff_mch=ff['apl_mch_oct']['sign']!=0 and ff['apl_mch_oct']['sign']==ff['pn_mch_oct']['sign']
    ff_dl=ff['apl_dl_oct']['sign']!=0 and ff['apl_dl_oct']['sign']==ff['pn_dl_oct']['sign']
    ff_pass=bool(ff_mch and ff_dl)
    g=identify_g(paths['g_delta'])
    lam=identify_lambda(paths['locality_primary'],paths['locality_secondary'],paths['mg_profiles'])
    gp=g['status']=='PASS_G_IDENTIFIED'; lp=lam['status']=='PASS_LAMBDA_IDENTIFIED'
    if not ff_pass: agg='FAIL_M2B_FEEDFORWARD_DIRECTIONAL_INCONSISTENCY'
    elif gp and lp: agg='PASS_M2B_G_AND_LAMBDA_IDENTIFIED_ETA_UNRESOLVED'
    elif gp: agg='PARTIAL_M2B_G_IDENTIFIED_LAMBDA_UNDERIDENTIFIED_ETA_UNRESOLVED'
    elif lp: agg='PARTIAL_M2B_LAMBDA_IDENTIFIED_G_UNDERIDENTIFIED_ETA_UNRESOLVED'
    else: agg='FAIL_M2B_PRIMARY_PARAMETERS_UNDERIDENTIFIED'
    report={'gate':'V6-M2B_PRISCO_TWO_PARAMETER_IDENTIFICATION','classification':agg,'dryad':{'doi':DOI,'version_id':version_id,'version_number':rec.get('versionNumber')},'manifest':manifest,'feedforward_directional':{'status':'PASS_FEEDFORWARD_DIRECTIONAL_CONSISTENCY' if ff_pass else 'FAIL_FEEDFORWARD_DIRECTIONAL_CONSISTENCY','details':ff},'g':g,'lambda':lam,'eta_ff':{'status':'UNRESOLVED_NOT_FITTED','value':None},'guardrails':{'amin_loaded':False,'mnq_market_loaded':False,'v5_residuals_used':False,'cross_cell_absolute_gcamp_ratio_used':False,'source_internal_distribution_files_used_for_fit':False}}
    p=OUT/'v6_m2b_prisco_identification.json'; p.write_text(json.dumps(report,indent=2,sort_keys=True)+'\n')
    print(json.dumps(report,indent=2,sort_keys=True))

if __name__=='__main__': main()
