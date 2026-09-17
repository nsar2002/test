from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

WORK = Path('v6_m1a_work')
OUT = Path('v6_m1a_out')
OUT.mkdir(exist_ok=True)

FILES = {
    'figure7': WORK / 'elife-56954-fig7-data1-v2.xlsx',
    'supplementary_stats': WORK / 'elife-56954-supp2-v2.xlsx',
}

_NUMERIC_LIKE_RE = re.compile(
    r'^\s*[<>~=]?\s*[-+]?(?:(?:\d+(?:[.,]\d*)?)|(?:[.,]\d+))(?:[eE][-+]?\d+)?\s*%?\s*$'
)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def normalize_string(value: str) -> str:
    # Remove invisible formatting/control characters before classification.
    value = ''.join(ch for ch in value if not unicodedata.category(ch).startswith('C'))
    return ' '.join(value.strip().split())


def is_numeric_like_string(value: str) -> bool:
    return bool(_NUMERIC_LIKE_RE.fullmatch(normalize_string(value)))


def schema_only(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        labels = []
        formula_cells = 0
        numeric_cells = 0
        numeric_like_string_cells = 0
        blank_cells = 0
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    blank_cells += 1
                    continue
                if cell.data_type == 'f':
                    formula_cells += 1
                    continue
                if isinstance(v, (int, float)):
                    numeric_cells += 1
                    continue
                if isinstance(v, str):
                    s = normalize_string(v)
                    if not s:
                        continue
                    if is_numeric_like_string(s):
                        numeric_like_string_cells += 1
                        continue
                    labels.append({'cell': cell.coordinate, 'text': s[:240]})
        text_blob = ' '.join(x['text'].lower() for x in labels)
        sheets.append({
            'title': ws.title,
            'max_row': int(ws.max_row or 0),
            'max_column': int(ws.max_column or 0),
            'string_labels': labels,
            'numeric_cell_count_only': int(numeric_cells),
            'numeric_like_string_cell_count_only': int(numeric_like_string_cells),
            'formula_cell_count_only': int(formula_cells),
            'blank_cell_count_only': int(blank_cells),
            'schema_tokens': {
                'mentions_calyx': 'calyx' in text_blob,
                'mentions_atp': 'atp' in text_blob,
                'mentions_kc_or_kenyon': ('kc' in text_blob) or ('kenyon' in text_blob),
                'mentions_distance_or_um': ('distance' in text_blob) or ('µm' in text_blob) or ('um' in text_blob),
                'mentions_normalized_or_inhib': ('normal' in text_blob) or ('inhib' in text_blob),
                'mentions_fly_or_neuron_or_n': ('fly' in text_blob) or ('neuron' in text_blob) or bool(re.search(r'\bn\b', text_blob)),
            },
        })
    return {'sheet_count': len(sheets), 'sheets': sheets}


def main() -> None:
    for p in FILES.values():
        if not p.exists():
            raise RuntimeError(f'missing source file {p}')

    schemas = {k: schema_only(p) for k, p in FILES.items()}
    report = {
        'gate': 'V6-M1A_AMIN_HOLDOUT_SOURCE_SCHEMA_PROBE',
        'status': 'SCHEMA_ACQUIRED_NOT_MODEL_COMPARED',
        'source': {
            k: {'path': str(p), 'sha256': sha256(p), 'bytes': p.stat().st_size}
            for k, p in FILES.items()
        },
        'schemas': schemas,
        'guardrails': {
            'response_numeric_cells_emitted': False,
            'numeric_like_string_values_emitted': False,
            'numeric_like_string_values_redacted': True,
            'numeric_values_used_for_parameter_selection': False,
            'v6_prediction_loaded': False,
            'mnq_or_market_loaded': False,
        },
        'correction_history': {
            'initial_probe_exposed_supplementary_pvalue_strings': True,
            'initial_probe_used_pvalues_for_parameter_or_threshold_selection': False,
            'canonical_probe_redacts_numeric_like_strings': True,
        },
    }
    (OUT / 'v6_m1a_amin_schema_probe.json').write_text(json.dumps(report, indent=2, sort_keys=True) + '\n')
    compact = {
        'status': report['status'],
        'source': report['source'],
        'guardrails': report['guardrails'],
        'correction_history': report['correction_history'],
        'sheets': {
            k: [
                {
                    'title': s['title'],
                    'max_row': s['max_row'],
                    'max_column': s['max_column'],
                    'numeric_cell_count_only': s['numeric_cell_count_only'],
                    'numeric_like_string_cell_count_only': s['numeric_like_string_cell_count_only'],
                    'schema_tokens': s['schema_tokens'],
                    'string_labels': s['string_labels'],
                }
                for s in v['sheets']
            ]
            for k, v in report['schemas'].items()
        },
    }
    print(json.dumps(compact, indent=2, sort_keys=True))


if __name__ == '__main__':
    main()
