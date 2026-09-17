from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: schema_probe.py DATASET.xlsx OUT.json")
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    wb = load_workbook(src, data_only=False, read_only=False)
    report = {
        "source_file": src.name,
        "sha256": sha256(src),
        "size_bytes": src.stat().st_size,
        "sheet_names": wb.sheetnames,
        "sheets": {},
        "numeric_values_exposed": False,
    }
    for ws in wb.worksheets:
        text_cells = []
        formula_cells = []
        numeric_count = 0
        nonempty_count = 0
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                nonempty_count += 1
                if isinstance(v, str):
                    if v.startswith("="):
                        formula_cells.append({"cell": cell.coordinate, "formula": v})
                    else:
                        text_cells.append({"cell": cell.coordinate, "text": v})
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    numeric_count += 1
        report["sheets"][ws.title] = {
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "nonempty_cells": nonempty_count,
            "numeric_cell_count": numeric_count,
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
            "text_cells": text_cells,
            "formula_cells": formula_cells,
        }
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
