from __future__ import annotations

import hashlib
import json
import re
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

DOI = 'doi:10.5061/dryad.bk3j9kdd1'
API = 'https://datadryad.org/api/v2'
DOWNLOAD_BASE = 'https://datadryad.org/stash/downloads/file_stream/'
WORK = Path('v6_m2a_work')
OUT = Path('v6_m2a_out')
WORK.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)

TARGETS = [
    'APL_PN_connectivity_table_(Figure_1D).xlsx',
    'APL_KC_connectivity_table_(Figure_1E).xlsx',
    'APL_GCaMP6m_Mch_vs_Oct_peaks_(Figure_2D).xlsx',
    'APL_GCaMP6m_DL_vs_Oct_peaks_(Figure_2G).xlsx',
    'GH146-GCaMP6m_Mch_vs_Oct_AL_average_activity_(Figure_3-figure_supplement_2B).xlsx',
    'GH146-GCaMP6m_DL_vs_Oct_AL_average_activity_(Figure_3-figure_supplement_2D).xlsx',
    'MB247-homer-GCaMP3_Mch_vs_Oct_average_response_(Figure_3F-H).xlsx',
    'APL-TNT_MB247-homer-GCaMP3_Mch_vs_Oct_(Figure_4B_and_4D).xlsx',
    'Inter-odour_delta_among_conditions_-_Figure_4-figure_supplement_1.xlsx',
    'APL_locality_PA_MP_vs_FA_MP_(Figure_5C).xlsx',
    'APL_locality_PA_FA_vs_MP_FA_(Figure_5-figure_supplement_1B).xlsx',
    'MB247-homer-GCaMP3_locality_of_MGs_responses_(Figure_5-figure_supplement_1A).xlsx',
    'NP225_Syp-GCaMP3_Mch_vs_Oct_all_boutons_peaks_(Figure_3C).xlsx',
    'MB247-homer-GCaMP3_Mch_vs_Oct_all_MG_peaks_(Figure_3G).xlsx',
    'APL-TNT_MB247-homer-GCaMP3_Mch_vs_Oct_all_MG_peaks_(Figure_4E_and_4F).xlsx',
]

_NUM_TOKEN = re.compile(r'(?<![A-Za-z])[-+]?(?:\d+(?:[.,]\d*)?|[.,]\d+)(?:[eE][-+]?\d+)?%?')
_PLIKE = re.compile(r'(?i)\b(?:p|r|r\^?2|mean|median|sem|sd)\s*[=<>]')


def request_json(url: str) -> dict:
    req = urllib.request.Request(url, headers={'Accept': 'application/json', 'X-API-Version': '2.1.0', 'User-Agent': 'v6-m2a-source-audit/1.0'})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.load(r)


def download(url: str, path: Path) -> None:
    req = urllib.request.Request(url, headers={'User-Agent': 'v6-m2a-source-audit/1.0'})
    with urllib.request.urlopen(req, timeout=120) as r, path.open('wb') as f:
        while True:
            chunk = r.read(1 << 20)
            if not chunk:
                break
            f.write(chunk)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def embedded_list(obj: dict, preferred: str) -> list[dict]:
    emb = obj.get('_embedded', {})
    if preferred in emb and isinstance(emb[preferred], list):
        return emb[preferred]
    for v in emb.values():
        if isinstance(v, list) and (not v or isinstance(v[0], dict)):
            return v
    return []


def resolve_version() -> dict:
    encoded = urllib.parse.quote(DOI, safe='')
    obj = request_json(f'{API}/datasets/{encoded}')
    link = obj.get('_links', {}).get('stash:version')
    href = link.get('href') if isinstance(link, dict) else None
    if not href:
        raise RuntimeError(f'Dryad dataset record lacks stash:version link; keys={list(obj)}')
    m = re.search(r'/versions/(\d+)(?:$|[/?#])', href)
    if not m:
        raise RuntimeError(f'cannot parse Dryad version id from href={href!r}')
    return {
        'id': int(m.group(1)),
        'versionNumber': obj.get('versionNumber'),
        'versionStatus': obj.get('versionStatus'),
        'lastModificationDate': obj.get('lastModificationDate'),
        'identifier': obj.get('identifier'),
    }


def resolve_files(version_id: int) -> list[dict]:
    url = f'{API}/versions/{version_id}/files?per_page=100'
    all_files: list[dict] = []
    seen = set()
    while url:
        if url in seen:
            raise RuntimeError('Dryad pagination loop')
        seen.add(url)
        obj = request_json(url)
        all_files.extend(embedded_list(obj, 'stash:files'))
        nxt = obj.get('_links', {}).get('next')
        href = nxt.get('href') if isinstance(nxt, dict) else None
        if not href:
            break
        url = href if href.startswith('http') else 'https://datadryad.org' + href
    return all_files


def file_name(meta: dict) -> str:
    return str(meta.get('path') or meta.get('name') or meta.get('fileName') or '').split('/')[-1]


def file_id(meta: dict) -> int:
    # Dryad v2 file list objects expose the file identifier in HAL links, not as a
    # top-level `id`. Parse the canonical /api/v2/files/{id} self link.
    links = meta.get('_links', {})
    for rel in ('self', 'stash:download'):
        link = links.get(rel)
        href = link.get('href') if isinstance(link, dict) else None
        if href:
            m = re.search(r'/files/(\d+)(?:/download)?(?:$|[/?#])', href)
            if m:
                return int(m.group(1))
    raise RuntimeError(f'cannot resolve Dryad file id from HAL links; rels={sorted(links)}')


def redact_label(s: str) -> tuple[str, bool]:
    s = ' '.join(str(s).strip().split())
    numeric_sensitive = bool(_NUM_TOKEN.search(s) or _PLIKE.search(s))
    if numeric_sensitive:
        s = _NUM_TOKEN.sub('[NUM]', s)
    return s[:240], numeric_sensitive


def workbook_schema(path: Path) -> dict:
    wb = load_workbook(path, read_only=False, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        labels = []
        counts = {'numeric': 0, 'formula': 0, 'blank': 0, 'boolean': 0, 'date': 0, 'numeric_like_or_numeric_token_string': 0, 'other': 0}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    counts['blank'] += 1
                elif c.data_type == 'f':
                    counts['formula'] += 1
                elif isinstance(v, bool):
                    counts['boolean'] += 1
                elif c.is_date:
                    counts['date'] += 1
                elif isinstance(v, (int, float)):
                    counts['numeric'] += 1
                elif isinstance(v, str):
                    txt, sensitive = redact_label(v)
                    if sensitive:
                        counts['numeric_like_or_numeric_token_string'] += 1
                    if txt:
                        labels.append({'cell': c.coordinate, 'text_redacted': txt})
                else:
                    counts['other'] += 1
        sheets.append({
            'title': ws.title,
            'max_row': int(ws.max_row or 0),
            'max_column': int(ws.max_column or 0),
            'merged_ranges': [str(x) for x in ws.merged_cells.ranges],
            'cell_type_counts': counts,
            'labels_with_numeric_tokens_redacted': labels,
        })
    return {'sheet_count': len(sheets), 'sheets': sheets}


def main() -> None:
    version = resolve_version()
    version_id = int(version['id'])
    files = resolve_files(version_id)
    by_name: dict[str, list[dict]] = {}
    for f in files:
        by_name.setdefault(file_name(f), []).append(f)

    missing = [n for n in TARGETS if n not in by_name]
    duplicates = {n: len(by_name[n]) for n in TARGETS if len(by_name.get(n, [])) != 1 and n in by_name}
    if missing or duplicates:
        raise RuntimeError(f'exact-file resolution failed missing={missing} duplicates={duplicates}')

    manifest = []
    schemas = {}
    for name in TARGETS:
        meta = by_name[name][0]
        fid = file_id(meta)
        dest = WORK / name
        download(f'{DOWNLOAD_BASE}{fid}', dest)
        local_hash = sha256(dest)
        provider_digest = meta.get('digest')
        provider_digest_type = meta.get('digestType')
        if str(provider_digest_type).lower() in {'sha-256', 'sha256'} and provider_digest and local_hash.lower() != str(provider_digest).lower():
            raise RuntimeError(f'provider SHA-256 mismatch for {name}: provider={provider_digest} local={local_hash}')
        if meta.get('size') is not None and int(meta['size']) != dest.stat().st_size:
            raise RuntimeError(f'provider size mismatch for {name}: provider={meta["size"]} local={dest.stat().st_size}')
        manifest.append({
            'filename': name,
            'dryad_file_id': fid,
            'provider_size': meta.get('size'),
            'provider_digest': provider_digest,
            'provider_digest_type': provider_digest_type,
            'local_bytes': dest.stat().st_size,
            'local_sha256': local_hash,
        })
        schemas[name] = workbook_schema(dest)

    roles = {
        'structural_context': TARGETS[0:2],
        'pn_to_apl_relative_recruitment_candidates': TARGETS[2:6],
        'causal_apl_output_candidates': TARGETS[6:9],
        'spatial_locality_candidates': TARGETS[9:12],
        'source_internal_falsification_only': TARGETS[12:15],
    }
    report = {
        'gate': 'V6-M2A_PRISCO_CALIBRATION_SOURCE_SCHEMA',
        'status': 'PASS_M2A_SOURCE_SCHEMA_ONLY',
        'dryad': {
            'doi': DOI,
            'version_id': version_id,
            'version_number': version.get('versionNumber'),
            'version_status': version.get('versionStatus'),
            'last_modification_date': version.get('lastModificationDate'),
            'total_files_seen': len(files),
        },
        'manifest': manifest,
        'roles': roles,
        'schemas': schemas,
        'guardrails': {
            'response_numeric_values_emitted': False,
            'numeric_tokens_in_string_labels_redacted': True,
            'parameters_fit': False,
            'v6_prediction_loaded': False,
            'amin_response_values_loaded': False,
            'mnq_or_market_loaded': False,
        },
    }
    out = OUT / 'v6_m2a_prisco_schema_probe.json'
    out.write_text(json.dumps(report, indent=2, sort_keys=True) + '\n')
    print(json.dumps({
        'status': report['status'],
        'dryad': report['dryad'],
        'manifest': manifest,
        'roles': roles,
        'guardrails': report['guardrails'],
        'schema_summary': {
            name: [
                {
                    'title': s['title'], 'max_row': s['max_row'], 'max_column': s['max_column'],
                    'merged_ranges': s['merged_ranges'], 'cell_type_counts': s['cell_type_counts'],
                    'labels_with_numeric_tokens_redacted': s['labels_with_numeric_tokens_redacted'],
                }
                for s in schema['sheets']
            ]
            for name, schema in schemas.items()
        },
    }, indent=2, sort_keys=True))


if __name__ == '__main__':
    main()
